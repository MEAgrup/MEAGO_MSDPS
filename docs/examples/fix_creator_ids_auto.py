#!/usr/bin/env python3
"""
Auto-fix Creator IDs: Add @ symbol to all Creator IDs in Creator Analysis format
Usage: python3 fix_creator_ids_auto.py <input_file> [output_file]
"""
import openpyxl
import sys
from pathlib import Path

def fix_creator_ids(input_path, output_path=None):
    """
    Fix Creator IDs by adding @ symbol if missing

    Args:
        input_path: Path to input XLSX file
        output_path: Path to output XLSX file (default: input_FIXED.xlsx)

    Returns:
        dict with status and stats
    """
    input_file = Path(input_path)

    if not input_file.exists():
        return {"ok": False, "error": f"File not found: {input_path}"}

    if not input_file.suffix.lower() == ".xlsx":
        return {"ok": False, "error": "File must be .xlsx format"}

    try:
        # Load workbook
        print(f"📂 Loading {input_file.name}...")
        wb = openpyxl.load_workbook(input_file)

        # Check sheets
        if "Filter" not in wb.sheetnames or "Data" not in wb.sheetnames:
            return {"ok": False, "error": "File must have 'Filter' and 'Data' sheets"}

        ws_data = wb["Data"]

        # Get headers to find Creator ID column
        headers = [cell.value for cell in ws_data[1]]
        if "Creator ID" not in headers:
            return {"ok": False, "error": "Column 'Creator ID' not found in Data sheet"}

        creator_id_col = headers.index("Creator ID") + 1  # Excel columns are 1-indexed

        print(f"📊 Sheet: {ws_data.title}")
        print(f"📋 Found {len(headers)} columns")
        print(f"🔍 Creator ID column: {chr(64 + creator_id_col)} (column {creator_id_col})")

        # Process data rows
        fixed_count = 0
        error_rows = []
        total_rows = 0

        for row_idx, row in enumerate(ws_data.iter_rows(min_row=2, values_only=False), start=2):
            total_rows += 1

            # Get cell for Creator ID
            cell = row[creator_id_col - 1]  # Convert to 0-indexed

            if cell.value is None:
                continue

            value = str(cell.value).strip()

            if not value:
                continue

            # Add @ if missing
            if not value.startswith("@"):
                try:
                    cell.value = "@" + value
                    fixed_count += 1
                except Exception as e:
                    error_rows.append({"row": row_idx, "error": str(e)})

        # Determine output file
        if output_path is None:
            stem = input_file.stem
            output_file = input_file.parent / f"{stem}_FIXED{input_file.suffix}"
        else:
            output_file = Path(output_path)

        # Save
        print(f"\n✅ Fixed {fixed_count}/{total_rows} Creator IDs (added @ symbol)")
        if error_rows:
            print(f"⚠️  Errors in {len(error_rows)} rows:")
            for err in error_rows[:5]:
                print(f"   Row {err['row']}: {err['error']}")

        print(f"💾 Saving to {output_file.name}...")
        wb.save(output_file)

        return {
            "ok": True,
            "input_file": str(input_file),
            "output_file": str(output_file),
            "total_rows": total_rows,
            "fixed_count": fixed_count,
            "error_count": len(error_rows),
            "message": f"✅ Successfully fixed {fixed_count} Creator IDs"
        }

    except Exception as e:
        return {
            "ok": False,
            "error": f"Error processing file: {str(e)}"
        }

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 fix_creator_ids_auto.py <input_file> [output_file]")
        print("\nExample:")
        print("  python3 fix_creator_ids_auto.py CreatorAnalysis_All_ManagedCreators_20260701_20260707.xlsx")
        print("  python3 fix_creator_ids_auto.py input.xlsx output.xlsx")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else None

    result = fix_creator_ids(input_file, output_file)

    print("\n" + "="*70)
    if result["ok"]:
        print("🎉 SUCCESS!")
        print("="*70)
        print(f"Input:  {result['input_file']}")
        print(f"Output: {result['output_file']}")
        print(f"\n📊 Statistics:")
        print(f"  Total rows processed: {result['total_rows']}")
        print(f"  Creator IDs fixed: {result['fixed_count']}")
        print(f"  Errors: {result['error_count']}")
        print(f"\n📋 Next steps:")
        print(f"  1. Upload {Path(result['output_file']).name}")
        print(f"  2. ✅ CENTANG 'Proses ulang jika file duplikat'")
        print(f"  3. Click 'Upload & Process'")
    else:
        print("❌ ERROR")
        print("="*70)
        print(f"Error: {result['error']}")
        sys.exit(1)

if __name__ == "__main__":
    main()
